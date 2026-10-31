"""Fill the official SDC finalist templates from our markdown sources, keeping the template's layout.

    python scripts/fill_submission.py <final-template.docx> <declaration-template.docx> <out-dir>

Sources: docs/submission/final-report.md and docs/submission/tech-declaration-final.md.
Needs python-docx (pip install python-docx). Writes SDC2026_HokkienMeeIsPurple_Final.docx and
SDC2026_HokkienMeeIsPurple_TechDeclaration_Final.docx; export them to PDF with Word.
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import docx
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
TICK, BOX = "☒", "☐"


# ---------------------------------------------------------------- markdown

def section(md: str, start: str, end: str | None = None) -> str:
    """Text after `start`, up to `end` or the next heading of the same or a higher level."""
    i = md.index(start) + len(start)
    if end:
        return md[i:md.index(end, i)].strip()
    level = start.split(" ")[0]
    m = re.search(rf"\n#{{1,{len(level)}}} ", md[i:])
    return md[i:i + m.start()].strip() if m else md[i:].strip()


def cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def table_rows(text: str) -> list[list[str]]:
    """Body rows of the first markdown table in `text` (no header, no --- line)."""
    rows = [cells(ln) for ln in text.splitlines() if ln.startswith("|")]
    return [r for r in rows[1:] if not set("".join(r)) <= set("-: ")]


def blocks(text: str) -> list[tuple[str, str]]:
    """("p" | "bullet" | "code", text) from markdown prose. A table becomes one line per row."""
    out, code = [], None
    for line in text.splitlines():
        if line.startswith("```"):
            if code is None:
                code = []
            else:
                out.append(("code", "\n".join(code)))
                code = None
        elif code is not None:
            code.append(line)
        elif line.startswith("|"):
            row = cells(line)
            if not set("".join(row)) <= set("-: "):
                out.append(("p", " · ".join(c for c in row if c)))
        elif not line.strip() or line.startswith("<!--"):
            continue
        elif line.startswith("- "):
            out.append(("bullet", line[2:]))
        else:
            out.append(("p", line))
    return out


# ---------------------------------------------------------------- docx

def runs(par, text: str, size: float | None = None) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


def clear(par) -> None:
    for r in list(par.runs):
        r._r.getparent().remove(r._r)


def fill_cell(cell, content, size: float = 10) -> None:
    if isinstance(content, str):
        content = [("p", line) for line in content.split("\n")]
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    first = cell.paragraphs[0]
    clear(first)
    for i, (kind, text) in enumerate(content):
        par = first if i == 0 else cell.add_paragraph()
        if kind == "code":
            for j, line in enumerate(text.split("\n")):
                lp = par if j == 0 else cell.add_paragraph()
                r = lp.add_run(line)
                r.font.name, r.font.size = "Consolas", Pt(7.5)
                lp.paragraph_format.space_after = Pt(0)
            continue
        runs(par, ("• " if kind == "bullet" else "") + text, size)
        par.paragraph_format.space_after = Pt(3)


def fit_rows(table, data: list[list[str]], size: float = 9) -> None:
    """Give the table exactly len(data) rows under its header row, then fill them."""
    while len(table.rows) - 1 < len(data):
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    for row in list(table.rows)[1 + len(data):]:
        table._tbl.remove(row._tr)
    for row, values in zip(table.rows[1:], data):
        for cell, value in zip(row.cells, values):
            fill_cell(cell, value, size)


def tick(par, choice: str) -> None:
    """'☐ Yes ☐ No ...' -> tick the box in front of `choice`."""
    text = par.text.replace(f"{BOX} {choice}", f"{TICK} {choice}", 1)
    clear(par)
    par.add_run(text)


def para_after(par, text: str) -> None:
    new = copy.deepcopy(par._p)
    par._p.addnext(new)
    p = docx.text.paragraph.Paragraph(new, par._parent)
    clear(p)
    runs(p, text)


def find_par(d, starts: str):
    return next(p for p in d.paragraphs if p.text.strip().startswith(starts))


def checklist(table, md_text: str) -> None:
    ticked = [ln.startswith("- [x]") for ln in md_text.splitlines() if ln.startswith("- [")]
    for row, on in zip(table.rows, ticked):
        fill_cell(row.cells[0], TICK if on else BOX)


def sign(d, md: str, signatures: list[tuple[str, str]] | None = None) -> None:
    """Fill [Signature] and [Date]. `signatures` = [(name, image path)], leader first; without it the line
    says the signed copy is separate (the repository is public, so signature images stay out of it)."""
    line = next(ln for ln in md.splitlines() if ln.startswith("Signature:"))
    _, name, date = [x.split(":", 1)[1].strip() for x in line.split(" · ")]
    for p in d.paragraphs:
        if p.text.strip() == "[Signature]":
            clear(p)
            if not signatures:
                p.add_run(f"Signature: signed copy submitted to the organisers    Name: {name} (team leader)")
                continue
            for i, (who, image) in enumerate(signatures):
                if i:
                    p.add_run().add_break()
                p.add_run("Team leader: " if i == 0 else "Team member: ")
                p.add_run().add_picture(image, height=Inches(0.55))
                p.add_run(f"  {who}")
        elif p.text.strip() == "[Date]":
            clear(p)
            p.add_run(f"Date: {date}")


# ---------------------------------------------------------------- the two documents

PHONE: str | None = None  # --phone: only for the copy sent to the organisers; the public repository never holds it


def private(md: str) -> str:
    return md.replace("phone on the submitted copy", PHONE) if PHONE else md


def final_report(template: Path, out: Path, signatures=None) -> None:
    md = private((ROOT / "docs/submission/final-report.md").read_text(encoding="utf-8"))
    d = docx.Document(str(template))
    t = d.tables  # the template: overview rows are tables 1-6, checklist rows are tables 13-17
    for table, (_, value) in zip(t[1:7], table_rows(section(md, "## 1. Team & Project Overview"))):
        fill_cell(table.cell(0, 1), value)
    fill_cell(t[7].cell(0, 0), blocks(section(md, "## 2. Project Description")))
    fill_cell(t[8].cell(0, 0), blocks(section(md, "### 3.1 Final System Architecture")))
    fill_cell(t[9].cell(0, 0), blocks(section(md, "### 3.2 Implementation & Hackathon Deliverables")))
    for row, (_, value) in zip(t[10].rows[1:], table_rows(section(md, "## 4. Environmental & SDG Impact Verification", "### 4.1"))):
        fill_cell(row.cells[1], value, 9)
    fill_cell(t[11].cell(0, 0), blocks(section(md, "### 4.1 Mathematical Proof / Code Logic Snippet")), 9)
    fill_cell(t[12].cell(0, 0), blocks(section(md, "## 5. Market Viability & Future Scalability Plan")))
    ticked = [ln.startswith("- [x]") for ln in section(md, "## 6. Final Submission Checklist").splitlines() if ln.startswith("- [")]
    for table, on in zip(t[13:18], ticked):
        box = table.cell(0, 0).tables[0].cell(0, 0)  # the drawn box is a nested table; a cell must still end in a paragraph
        fill_cell(box, "X" if on else "")

    boxes = [p for p in d.paragraphs if p.text.strip().startswith(f"{BOX} Yes") and "No change" in p.text]
    tick(boxes[0], "No change from Phase 1 declaration")  # 7.1
    tick(boxes[1], "Yes")                                 # 7.2
    if section(md, "### 7.3 Employment-Related IP Conflicts").startswith("No."):
        tick(boxes[2], "No")                              # 7.3
    para_after(find_par(d, "If yes, briefly describe what was pre-existing"), section(md, "### 7.1 Pre-Existing / Prior-Owned Work"))
    para_after(find_par(d, "If yes, list them below"), section(md, "### 7.2 Third-Party Code, Datasets, or Models").splitlines()[0])
    fit_rows(t[18], table_rows(section(md, "### 7.2 Third-Party Code, Datasets, or Models")))
    para_after(find_par(d, "If yes, please describe the nature of the agreement"), section(md, "### 7.3 Employment-Related IP Conflicts"))
    sign(d, md, signatures)
    d.save(str(out))


def declaration(template: Path, out: Path, signatures=None) -> None:
    md = private((ROOT / "docs/submission/tech-declaration-final.md").read_text(encoding="utf-8"))
    d = docx.Document(str(template))
    t = d.tables
    for row, (_, value) in zip(t[1].rows, table_rows(section(md, "## 1. Team & Project Overview"))):
        fill_cell(row.cells[1], value)
    fill_cell(t[2].cell(0, 0), blocks(section(md, "### 2.1 System Architecture Overview")))
    fit_rows(t[3], table_rows(section(md, "### 2.2 Languages, Frameworks & Libraries")))
    fit_rows(t[4], table_rows(section(md, "### 3.1 External APIs")))
    fit_rows(t[5], table_rows(section(md, "### 3.2 Cloud Platforms & Deployment")))
    models = section(md, "### 4.1 AI Models Used in the Solution")
    fit_rows(t[6], table_rows(models))
    para_after(find_par(d, "List every AI model embedded"), models.splitlines()[-1])
    fit_rows(t[7], table_rows(section(md, "### 4.2 Datasets")))
    yes_no = [p for p in d.paragraphs if p.text.strip().startswith(f"{BOX} Yes") and "No change" not in p.text]
    for p in yes_no:  # 4.3 and 5: both Yes
        tick(p, "Yes")
    fit_rows(t[8], table_rows(section(md, "### 4.3 AI-Assisted Development Tools")))
    fill_cell(t[9].cell(0, 0), section(md, "**How the team reviewed AI output.**", "## 5."))
    fit_rows(t[10], table_rows(section(md, "## 5. Sponsor-Provided Resources")))
    fill_cell(t[11].cell(0, 0), section(md, "**Essential or replaceable?**", "## 6."))
    fill_cell(t[12].cell(0, 0), section(md, "### 6.1 Data Handling & Privacy"))
    fill_cell(t[13].cell(0, 0), blocks(section(md, "### 6.2 Known Limitations & Risks")))
    checklist(t[14], section(md, "## 7. Compliance & Readiness Checklist"))
    sign(d, md, signatures)
    d.save(str(out))


if __name__ == "__main__":
    # --sign "Tay Ee Zhan=leader.png" "Goh Sheng Kuan=member.png" puts the signature images in; leader first.
    # --phone "+60 ..." fills the leader's phone (place it before --sign).
    args = sys.argv[1:]
    if "--phone" in args:
        PHONE = args[args.index("--phone") + 1]
    sigs = [tuple(x.split("=", 1)) for x in args[args.index("--sign") + 1:]] if "--sign" in args else None
    final_tpl, decl_tpl, out_dir = Path(args[0]), Path(args[1]), Path(args[2])
    final_report(final_tpl, out_dir / "SDC2026_HokkienMeeIsPurple_Final.docx", sigs)
    declaration(decl_tpl, out_dir / "SDC2026_HokkienMeeIsPurple_TechDeclaration_Final.docx", sigs)
    print("written to", out_dir)
